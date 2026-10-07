import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCREENSHOTS_DIR = os.path.join(SKILL_DIR, "screenshots")
# Fallback to scratch if skill screenshots not yet populated
FALLBACK_DIR = r"d:\dashboard-bvl-v2\scratch\manual_screenshots"
DOCS_DIR = r"d:\dashboard-bvl-v2\docs"

def get_image_path(filename):
    p1 = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(p1): return p1
    p2 = os.path.join(FALLBACK_DIR, filename)
    if os.path.exists(p2): return p2
    return None

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_step_card(doc, step_num, title, action_text, note=None):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    c0 = tbl.cell(0, 0)
    c0.width = Inches(0.8)
    set_cell_background(c0, "1E3A5F")
    set_cell_margins(c0, top=80, bottom=80, left=60, right=60)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(f"LANGKAH\n{step_num}")
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(8.5)
    r0.font.color.rgb = RGBColor(255, 255, 255)
    
    c1 = tbl.cell(0, 1)
    c1.width = Inches(5.6)
    set_cell_background(c1, "F8FAFC")
    set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    
    r_t = p1.add_run(f"{title}: ")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = RGBColor(30, 58, 95)
    
    r_a = p1.add_run(action_text)
    r_a.font.name = "Arial"
    r_a.font.size = Pt(9.5)
    r_a.font.color.rgb = RGBColor(51, 65, 85)
    
    if note:
        p_n = c1.add_paragraph()
        p_n.paragraph_format.space_before = Pt(2)
        p_n.paragraph_format.space_after = Pt(0)
        r_n = p_n.add_run(f"💡 Tips: {note}")
        r_n.font.name = "Arial"
        r_n.font.size = Pt(8.5)
        r_n.font.italic = True
        r_n.font.color.rgb = RGBColor(180, 83, 9)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_box(doc, img_filename, caption_text):
    img_path = get_image_path(img_filename)
    if not img_path:
        print(f"Warning: Image {img_filename} not found.")
        return
        
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    c_img = tbl.cell(0, 0)
    c_img.width = Inches(6.4)
    set_cell_background(c_img, "F1F5F9")
    set_cell_margins(c_img, top=60, bottom=40, left=60, right=60)
    p_img = c_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(2)
    
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.2))
    
    c_cap = tbl.cell(1, 0)
    c_cap.width = Inches(6.4)
    set_cell_background(c_cap, "E2E8F0")
    set_cell_margins(c_cap, top=40, bottom=40, left=80, right=80)
    p_cap = c_cap.paragraphs[0]
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(0)
    
    r_cap = p_cap.add_run(f"📸 Gambar: {caption_text}")
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(8.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(8)

def add_module_header(doc, number, title, icon, url, purpose, timing):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(f"{icon} MODUL {number}: {title.upper()}")
    r.font.name = "Georgia"
    r.font.size = Pt(13.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 95)

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [Inches(1.8), Inches(2.8), Inches(1.8)]
    
    c0 = tbl.cell(0, 0)
    c0.width = col_w[0]
    set_cell_background(c0, "EFF6FF")
    set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
    p0 = c0.paragraphs[0]
    r_u1 = p0.add_run("🌐 URL / Menu:\n")
    r_u1.bold = True
    r_u1.font.size = Pt(8)
    r_u1.font.color.rgb = RGBColor(37, 99, 235)
    r_u2 = p0.add_run(url)
    r_u2.font.size = Pt(8.5)
    r_u2.font.bold = True
    r_u2.font.color.rgb = RGBColor(15, 23, 42)

    c1 = tbl.cell(0, 1)
    c1.width = col_w[1]
    set_cell_background(c1, "F8FAFC")
    set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
    p1 = c1.paragraphs[0]
    r_p1 = p1.add_run("🎯 Tujuan Pengisian:\n")
    r_p1.bold = True
    r_p1.font.size = Pt(8)
    r_p1.font.color.rgb = RGBColor(71, 85, 105)
    r_p2 = p1.add_run(purpose)
    r_p2.font.size = Pt(8.5)
    r_p2.font.color.rgb = RGBColor(51, 65, 85)

    c2 = tbl.cell(0, 2)
    c2.width = col_w[2]
    set_cell_background(c2, "FEF3C7")
    set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
    p2 = c2.paragraphs[0]
    r_t1 = p2.add_run("⏰ Waktu Input:\n")
    r_t1.bold = True
    r_t1.font.size = Pt(8)
    r_t1.font.color.rgb = RGBColor(180, 83, 9)
    r_t2 = p2.add_run(timing)
    r_t2.font.size = Pt(8.5)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(146, 64, 14)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def build_visual_manual(output_filepath):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    navy = RGBColor(30, 58, 95)
    gold = RGBColor(197, 160, 89)
    slate = RGBColor(71, 85, 105)

    # TITLE BANNER
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hdr = p_hdr.add_run("BVLGARI INDONESIA — PT MOGEMS PUTRI INTERNATIONAL")
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(9.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = gold

    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main.paragraph_format.space_before = Pt(4)
    p_main.paragraph_format.space_after = Pt(2)
    r_main = p_main.add_run("PANDUAN PRAKTIS & BERGAMBAR\nAPA SAJA YANG HARUS DI-INPUT OLEH STORE MANAGER (SM)")
    r_main.font.name = "Georgia"
    r_main.font.size = Pt(17)
    r_main.font.bold = True
    r_main.font.color.rgb = navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Panduan Visual Step-by-Step Langkah 1-2-3 Khusus Butik Plaza Indonesia, Plaza Senayan & Bali\nPortal: https://bvl.mogems.co.id")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = slate

    # 1. SETUP & TARGETS
    add_module_header(doc, 1, "Setup Target Butik & Target Advisor", "🎯", "/advisor-setup", "Menentukan target omzet toko & target individu Client Advisor untuk bulan berjalan.", "Awal Bulan (Tgl 1) atau Awal Tahun")
    add_image_box(doc, "3_advisor_setup.png", "Layar Menu Setup & Targets (Tab Homebase, Rotation, Target Advisor, Store Target)")
    add_step_card(doc, "1", "Pilih Tab 'Store Target'", "Klik tab 'Store Target' di bagian atas. Masukkan target rupiah penjualan butik Anda untuk setiap bulan (Januari-Desember).", note="Angka target ini akan otomatis menjadi target di halaman Operations Sales dan Daily Report.")
    add_step_card(doc, "2", "Pilih Tab 'Target' (Individu)", "Pilih bulan berjalan. Masukkan target nominal penjualan untuk masing-masing Client Advisor Anda.", note="Store Manager / Supervisor diberi nilai 0 agar tidak masuk dalam perhitungan ranking advisor penjualan.")
    add_step_card(doc, "3", "Pilih Tab 'Rotation' (Jika Ada Rotasi)", "Jika ada advisor Anda yang diperbantukan sementara ke butik lain, pilih lokasi butik penugasan pada bulan terkait.", note="Sistem otomatis menghitung prorata pencapaian target advisor sesuai lokasi bertugas.")
    add_step_card(doc, "4", "Klik Tombol Simpan (Save)", "Setelah angka terisi, klik ikon centang hijau (✓) atau tombol Simpan di kanan baris.", note="Data tersimpan langsung ke database cloud dan angka ranking akan ter-update seketika.")

    # 2. SALES JOURNAL
    add_module_header(doc, 2, "Sales Journal (Buku Catatan Harian Toko)", "📔", "/sales-journal", "Mencatat peristiwa harian toko (alasan sales naik/turun, kendala EDC/POS, cuaca, promo mall, tamu VIP).", "Setiap Malam Saat Closing (21:45)")
    add_image_box(doc, "4_sales_journal.png", "Layar Sales Journal dengan Daftar Tanggal, Traffic, dan Kolom Catatan Butik")
    add_step_card(doc, "1", "Pilih Tanggal Hari Ini", "Cari baris tanggal hari ini pada daftar kalender, lalu klik baris tersebut untuk membuka kotak input catatan.")
    add_step_card(doc, "2", "Ketik Catatan Kondisi Toko", "Tuliskan peristiwa nyata yang terjadi di butik Anda hari ini. Contoh: Kunjungan VIP Client beli Serpenti, EDC BCA offline, hujan lebat mall sepi.", note="Catatan ini dibaca oleh Head of Operations dan Direksi untuk memahami kondisi riil toko.")
    add_step_card(doc, "3", "Pilih Label Tag Kategori", "Centang pilihan tag yang sesuai: Promo / Event / VIP Customer / Weather / Stockout / Staffing / Complaint.")
    add_step_card(doc, "4", "Klik 'Simpan Catatan'", "Tekan tombol Simpan Catatan. Catatan Anda akan tersimpan permanen dan dapat dirangkum otomatis oleh fitur AI Synthesis.")

    # 3. DAILY REPORT & EMAIL
    add_module_header(doc, 3, "Daily Report & Pengiriman Email Laporan Closing", "✉️", "/daily-report", "Memeriksa hasil penjualan hari ini dan mengirimkan email laporan resmi ke Direksi & Manajemen.", "Setiap Malam Saat Closing (22:00)")
    add_image_box(doc, "2_daily_report.png", "Layar Daily Sales Report dengan Tombol Download PDF, Excel, WhatsApp, dan Send Email")
    add_step_card(doc, "1", "Periksa Angka Closing Hari Ini", "Klik tombol 'Today' atau pastikan tanggal yang dipilih adalah hari ini. Periksa Total Penjualan Store dan MTD Achievement %.")
    add_step_card(doc, "2", "Klik Tombol 'Send Email'", "Klik tombol Send Email di pojok kanan atas bilah filter. Jendela popup mewah (Modal Live Preview) akan terbuka.")
    add_image_box(doc, "2b_daily_email_modal.png", "Modal Live Preview Email — Menampilkan Pratinjau Tampilan Outlook & Pengaturan Penerima")
    add_step_card(doc, "3", "Tinjau Tab 1: Live Preview", "Pastikan tata letak email, tabel crossing sales, dan breakdown harian sudah tampil rapi dan angka penjualan sesuai fisik kasir.")
    add_step_card(doc, "4", "Buka Tab 2: Penerima & Lampiran", "Pastikan alamat email sudah benar:\n• To: renaldi@mraretail.co.id (Pak Aldi)\n• CC: jessica@mogems.co.id, natalia@mraretail.co.id, aris@mraretail.co.id\n• Pastikan centang Berkas PDF dan Berkas Excel tetap aktif.")
    add_step_card(doc, "5", "Klik 'Kirim Email Sekarang'", "Klik tombol hitam 'Kirim Email Sekarang'. Tunggu beberapa detik hingga muncul notifikasi hijau 'Email Berhasil Dikirim'.", note="Laporan resmi PDF dan Excel langsung terlampir otomatis ke email Direksi.")

    # 4. CRM & TRAFFIC
    add_module_header(doc, 4, "CRM & Traffic (Input Tamu Walk-in & Wishlist)", "👥", "/crm-dedup", "Mencatat data tamu walk-in yang datang ke butik dan menautkan barang yang diminati (live inventory).", "Setiap Kali Ada Tamu Berkunjung")
    add_image_box(doc, "5_crm_traffic.png", "Layar CRM & Traffic untuk Input Tamu Baru, Pencarian Stok Produk, dan Audit Duplikasi")
    add_step_card(doc, "1", "Buka Tab 'Traffic'", "Klik tab Traffic di bagian atas layar untuk membuka formulir kunjungan tamu.")
    add_step_card(doc, "2", "Isi Data Tamu", "Masukkan Nama Lengkap tamu, Nomor HP / WhatsApp aktif, dan Kota/Negara domisili.")
    add_step_card(doc, "3", "Pilih Advisor Pendamping", "Pilih nama Client Advisor butik Anda yang mendampingi dan melayani tamu tersebut.")
    add_step_card(doc, "4", "Cari Barang Minat (Live Stok)", "Ketik nama barang yang dicoba atau diminati tamu (misal: 'Octo', 'Serpenti Tubogas', 'B.zero1'). Pilih item dari daftar inventori toko untuk disimpan sebagai wishlist klien.", note="Fitur ini memudahkan advisor melakukan follow-up saat barang atau promo baru tiba.")
    add_step_card(doc, "5", "Klik Simpan Kunjungan", "Data tamu otomatis masuk ke database CRM dan terhubung dengan riwayat traffic butik.")

    # 5. INVOICE MANAGEMENT
    add_module_header(doc, 5, "Invoice Management (Audit Kasir & Split-Payment)", "🧾", "/invoice-management", "Memeriksa struk POS OMEGA, mengoreksi salah pilih toko, dan verifikasi split-payment EDC.", "Setiap Malam Sebelum Closing (21:30)")
    add_image_box(doc, "6_invoice_management.png", "Layar Invoice Management — Rincian Invoice, Detail Pembayaran POS, dan Catatan Verifikasi")
    add_step_card(doc, "1", "Cek Daftar Invoice Hari Ini", "Pilih bulan dan tahun. Cari nomor invoice struk kasir yang ingin diverifikasi.")
    add_step_card(doc, "2", "Input Catatan Verifikasi (Ikon Catatan)", "Klik tombol catatan di sebelah nomor invoice. Masukkan keterangan audit kasir, misalnya: 'Split payment 2 kartu disetujui SM' atau 'Struk pengganti karena salah input nama advisor'.", note="Catatan ini dibaca oleh tim Finance untuk proses rekonsiliasi bank.")
    add_step_card(doc, "3", "Koreksi Lokasi Butik (Jika Kasir Salah)", "Jika kasir keliru memilih cabang butik di sistem OMEGA POS, klik dropdown lokasi dan ubah ke butik yang benar.")
    add_step_card(doc, "4", "Koreksi Komisi Kartu (Jika EDC Silang)", "Jika kartu bank digesek pada mesin EDC berbeda, SM dapat menyesuaikan persentase komisi bank pada rincian pembayaran.")

    # 6. INSTALLMENT & MDR GUIDE
    add_module_header(doc, 6, "Panduan Cicilan 0% & MDR Bank (Aturan Gesek EDC)", "💳", "/installment-guide", "Panduan resmi komisi mesin EDC dan matriks cicilan 0% 11 bank agar kasir tidak salah gesek.", "Referensi Wajib Setiap Ada Transaksi Kartu")
    add_image_box(doc, "7_installment_guide.png", "Layar Installment & MDR Guide — Matriks Cicilan 0% 11 Bank & Aturan Gesek EDC")
    add_step_card(doc, "1", "Buka Matriks Cicilan 11 Bank", "Pilih bank kartu kredit yang digunakan pelanggan. Periksa tenor cicilan yang berlaku (3, 6, 12, atau 24 bulan).")
    add_step_card(doc, "2", "Aturan Wajib: Gesek di Mesin EDC Bank yang Sama!", "PASTIKAN KARTU KREDIT DIGESEK DI MESIN EDC BANK YANG SAMA DENGAN PENERBIT KARTU (Contoh: Kartu Kredit BCA di EDC BCA, Kartu Kredit Mandiri di EDC Mandiri).\n⚠️ Gesek kartu bank lain di mesin berbeda (Cross-EDC) akan dikenakan potongan komisi reguler yang jauh lebih mahal bagi toko!", note="Ini adalah aturan wajib nomor satu untuk seluruh kasir butik Bvlgari.")
    add_step_card(doc, "3", "Periksa Kartu Khusus (AMEX, WeChat, Alipay)", "Untuk kartu American Express (AMEX) dan e-wallet turis asing (WeChat Pay & Alipay), gunakan terminal pembayaran khusus sesuai instruksi di halaman ini.")

    # 7. ADVISOR PERFORMANCE
    add_module_header(doc, 7, "Advisor Performance & Evaluasi Tim Penjualan", "🏆", "/advisor-performance", "Memantau peringkat penjualan advisor, ekspor rekap 6 bulan prorata, dan kirim laporan advisor.", "Setiap Hari (Review) & Akhir Bulan")
    add_image_box(doc, "8_advisor_performance.png", "Layar Advisor Performance — Peringkat Advisor, % Pencapaian Target, dan Ekspor Excel")
    add_step_card(doc, "1", "Pantau Pencapaian Target Advisor", "Periksa kolom Net Sales, Target, dan Achievement % setiap Client Advisor. Warna Bar:\n• 🔵 Biru: Di atas 100% (Target Tercapai)\n• 🟢 Hijau: 80% - 99.9% (On Track)\n• 🔴 Merah: Di bawah 80% (Perlu Dimotivasi Khusus)")
    add_step_card(doc, "2", "Ekspor Excel Rata-rata 6 Bulan", "Klik tombol 'Export Rata-rata 6 Bln'. Berkas Excel akan terunduh otomatis dengan perhitungan proporsional bagi advisor baru dan menyaring staff non-target.")
    add_step_card(doc, "3", "Kirim Email Laporan Advisor", "Klik tombol 'Send Email' di pojok kanan atas untuk membuka modal preview laporan bulanan advisor ke manajemen.")

    # SAVE
    os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
    doc.save(output_filepath)
    print(f"SUCCESS: Visual Word manual generated at {output_filepath}")

if __name__ == "__main__":
    out_file = os.path.join(DOCS_DIR, "PANDUAN_BERGAMBAR_STORE_MANAGER_BVLGARI.docx")
    build_visual_manual(out_file)
